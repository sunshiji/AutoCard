import os
import re
from abc import ABC, abstractmethod
from datetime import datetime
from typing import Dict, List, Optional, Any, Tuple
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from models import (
    ResumeData, PersonalInfo, EducationExperience,
    WorkExperience, ProjectExperience
)


class BaseParser(ABC):
    @abstractmethod
    def parse(self, file_path: str) -> ResumeData:
        pass
    
    def _parse_date(self, date_str: str) -> Optional[datetime]:
        if not date_str:
            return None
        
        date_patterns = [
            r'(\d{4})年(\d{1,2})月',
            r'(\d{4})-(\d{1,2})',
            r'(\d{4})/(\d{1,2})',
            r'(\d{4})\.(\d{1,2})',
        ]
        
        for pattern in date_patterns:
            match = re.search(pattern, date_str)
            if match:
                try:
                    year = int(match.group(1))
                    month = int(match.group(2))
                    return datetime(year, month, 1)
                except ValueError:
                    continue
        
        year_match = re.search(r'(\d{4})', date_str)
        if year_match:
            try:
                return datetime(int(year_match.group(1)), 1, 1)
            except ValueError:
                pass
        
        return None
    
    def _extract_phone(self, text: str) -> Optional[str]:
        phone_patterns = [
            r'1[3-9]\d{9}',
            r'\+?86[-\s]?1[3-9]\d{9}',
            r'\d{3}[-\s]?\d{4}[-\s]?\d{4}',
            r'手机[：:]\s*([1][3-9]\d{9})',
            r'电话[：:]\s*([1][3-9]\d{9})',
            r'联系电话[：:]\s*([1][3-9]\d{9})',
            r'mobile[：:\s]+([1][3-9]\d{9})',
            r'phone[：:\s]+([1][3-9]\d{9})',
        ]
        
        for pattern in phone_patterns:
            match = re.search(pattern, text)
            if match:
                phone = match.group()
                if match.lastindex and match.group(1):
                    phone = match.group(1)
                phone = phone.replace('-', '').replace(' ', '').replace('+', '')
                if phone.startswith('86'):
                    phone = phone[2:]
                if len(phone) == 11 and phone.startswith('1'):
                    return phone
        return None
    
    def _extract_email(self, text: str) -> Optional[str]:
        email_patterns = [
            r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',
            r'邮箱[：:]\s*([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})',
            r'email[：:\s]+([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})',
            r'mail[：:\s]+([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})',
        ]
        
        for pattern in email_patterns:
            match = re.search(pattern, text)
            if match:
                if match.lastindex and match.group(1):
                    return match.group(1).strip()
                return match.group()
        return None
    
    def _extract_name(self, text: str, lines: List[str]) -> Optional[str]:
        name_patterns = [
            r'姓名[：:]\s*([^\s，。,；;\n\r]+)',
            r'名字[：:]\s*([^\s，。,；;\n\r]+)',
            r'^([\u4e00-\u9fa5]{2,4})$',
            r'申请人[：:]\s*([^\s，。,；;\n\r]+)',
            r'求职人[：:]\s*([^\s，。,；;\n\r]+)',
            r'Name[：:\s]+([a-zA-Z\s]+)',
        ]
        
        for pattern in name_patterns:
            match = re.search(pattern, text)
            if match:
                name = match.group(1).strip()
                if re.match(r'^[\u4e00-\u9fa5]{2,4}$', name):
                    return name
                if re.match(r'^[a-zA-Z\s]+$', name):
                    return name
        
        if lines:
            first_line = lines[0].strip()
            if 2 <= len(first_line) <= 4 and re.match(r'^[\u4e00-\u9fa5]+$', first_line):
                return first_line
            
            if len(lines) >= 2:
                for line in lines[:3]:
                    line = line.strip()
                    if 2 <= len(line) <= 4 and re.match(r'^[\u4e00-\u9fa5]+$', line):
                        phone_match = re.search(r'1[3-9]\d{9}', line)
                        email_match = re.search(r'@', line)
                        if not phone_match and not email_match:
                            return line
        
        return None
    
    def _extract_gender(self, text: str) -> Optional[str]:
        gender_patterns = [
            r'性别[：:]\s*(男|女)',
            r'\b(男|女)\b',
            r'性别[：:]\s*(male|female)',
            r'性别[：:]\s*(先生|女士)',
            r'性别[：:]\s*(男|女|未知)',
        ]
        
        for pattern in gender_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            if matches:
                gender = matches[0]
                if gender.lower() in ['male', '先生']:
                    return '男'
                elif gender.lower() in ['female', '女士']:
                    return '女'
                return gender
        return None
    
    def _extract_education(self, text: str) -> Optional[str]:
        education_keywords = [
            '博士', '硕士', '本科', '大专', '高中', '中专',
            '博士研究生', '硕士研究生', '研究生',
            '学士', '专科', '初中', '小学',
            'PhD', 'Master', 'Bachelor', 'College',
        ]
        
        education_order = [
            '博士研究生', '硕士研究生', '博士', '硕士', 
            '研究生', '本科', '学士', '大专', '专科',
            '高中', '中专', '初中', '小学',
            'PhD', 'Master', 'Bachelor', 'College',
        ]
        
        found_educations = []
        for keyword in education_keywords:
            if keyword in text:
                found_educations.append(keyword)
        
        if found_educations:
            for edu in education_order:
                if edu in found_educations:
                    return edu
        
        return None
    
    def _split_by_sections(self, text: str) -> Dict[str, str]:
        section_patterns = [
            r'[\n\r]+([\u4e00-\u9fa5]{2,8})[：: 　]*[\n\r]+',
            r'【([\u4e00-\u9fa5]{2,8})】',
            r'■([\u4e00-\u9fa5]{2,8})',
            r'◆([\u4e00-\u9fa5]{2,8})',
            r'●([\u4e00-\u9fa5]{2,8})',
            r'★([\u4e00-\u9fa5]{2,8})',
            r'▶([\u4e00-\u9fa5]{2,8})',
            r'►([\u4e00-\u9fa5]{2,8})',
        ]
        
        sections = {}
        current_section = 'basic_info'
        sections[current_section] = ''
        
        lines = text.split('\n')
        for line in lines:
            is_section_header = False
            for pattern in section_patterns:
                match = re.match(pattern, line.strip())
                if match:
                    current_section = match.group(1).strip()
                    sections[current_section] = ''
                    is_section_header = True
                    break
            
            if not is_section_header:
                sections[current_section] += line + '\n'
        
        if len(sections) <= 1:
            sections = self._split_by_sections_fuzzy(text)
        
        return sections
    
    def _split_by_sections_fuzzy(self, text: str) -> Dict[str, str]:
        sections = {}
        current_section = 'basic_info'
        sections[current_section] = ''
        
        section_keywords = {
            '教育': '教育经历',
            '工作': '工作经历',
            '项目': '项目经历',
            '技能': '技能',
            '自我介绍': '自我介绍',
            '个人简介': '自我介绍',
            '自我评价': '自我介绍',
            '基本信息': '基本信息',
            '个人信息': '个人信息',
            '求职意向': '求职意向',
        }
        
        lines = text.split('\n')
        
        for line in lines:
            line_stripped = line.strip()
            
            if not line_stripped:
                sections[current_section] += line + '\n'
                continue
            
            matched = False
            for keyword, section_name in section_keywords.items():
                if keyword in line_stripped and len(line_stripped) < 15:
                    if re.search(rf'^[■◆●★▶►\s]*{keyword}', line_stripped):
                        current_section = section_name
                        sections[current_section] = ''
                        matched = True
                        break
            
            if not matched:
                sections[current_section] += line + '\n'
        
        return sections
    
    def _parse_education_section(self, section_text: str) -> List[EducationExperience]:
        experiences = []
        
        edu_blocks = self._split_experience_blocks(section_text)
        
        if not edu_blocks or len(edu_blocks) == 0:
            edu_blocks = self._split_education_fuzzy(section_text)
        
        for block in edu_blocks:
            exp = EducationExperience()
            
            lines = [l.strip() for l in block.split('\n') if l.strip()]
            
            date_range = self._extract_date_range(block)
            if date_range:
                exp.start_date = date_range[0]
                exp.end_date = date_range[1]
            
            school_patterns = [
                r'([\u4e00-\u9fa5]{2,}(大学|学院|学校|研究院|研究所|警校|军校))',
                r'University\s+of\s+[\w]+',
                r'[\w]+\s+University',
                r'[\w]+\s+College',
                r'学校[：:]\s*([^\n\r，。；;\s]+)',
                r'毕业院校[：:]\s*([^\n\r，。；;\s]+)',
                r'院校[：:]\s*([^\n\r，。；;\s]+)',
            ]
            
            for pattern in school_patterns:
                match = re.search(pattern, block)
                if match:
                    if match.lastindex:
                        exp.school = match.group(1).strip()
                    else:
                        exp.school = match.group().strip()
                    break
            
            major_patterns = [
                r'专业[：:]\s*([^\n\r，。；;\n]+)',
                r'主修[：:]\s*([^\n\r，。；;\n]+)',
                r'所学专业[：:]\s*([^\n\r，。；;\n]+)',
                r'专业方向[：:]\s*([^\n\r，。；;\n]+)',
                r'([\u4e00-\u9fa5]{2,10}专业)',
            ]
            
            for pattern in major_patterns:
                match = re.search(pattern, block)
                if match:
                    if match.lastindex:
                        exp.major = match.group(1).strip()
                    else:
                        exp.major = match.group().strip()
                    break
            
            exp.education = self._extract_education(block)
            
            if not exp.education:
                edu_indicators = [
                    ('博士研究生', '博士'),
                    ('硕士研究生', '硕士'),
                    ('研究生', '硕士'),
                    ('本科', '本科'),
                    ('学士', '本科'),
                    ('大专', '大专'),
                    ('专科', '大专'),
                    ('高中', '高中'),
                ]
                
                for indicator, edu in edu_indicators:
                    if indicator in block:
                        exp.education = edu
                        break
            
            if exp.school or exp.major or exp.education or exp.start_date:
                experiences.append(exp)
        
        return experiences
    
    def _split_education_fuzzy(self, text: str) -> List[str]:
        blocks = []
        
        school_keywords = [
            '大学', '学院', '学校', '研究院', '研究所',
            'University', 'College', 'university', 'college'
        ]
        
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        
        current_block = ''
        in_education = False
        
        for line in lines:
            has_school = any(kw in line for kw in school_keywords)
            has_date = re.search(r'\d{4}[\s年./-]*\d{0,2}', line)
            
            if has_school or has_date:
                if current_block:
                    blocks.append(current_block.strip())
                current_block = line + '\n'
                in_education = True
            elif in_education:
                current_block += line + '\n'
        
        if current_block:
            blocks.append(current_block.strip())
        
        if not blocks:
            blocks = [text.strip()]
        
        return blocks
    
    def _parse_work_section(self, section_text: str) -> List[WorkExperience]:
        experiences = []
        
        work_blocks = self._split_experience_blocks(section_text)
        
        if not work_blocks or len(work_blocks) == 0:
            work_blocks = self._split_work_fuzzy(section_text)
        
        for block in work_blocks:
            exp = WorkExperience()
            
            date_range = self._extract_date_range(block)
            if date_range:
                exp.start_date = date_range[0]
                exp.end_date = date_range[1]
            
            company_patterns = [
                r'公司[：:]\s*([^\n\r，。；;\s]+)',
                r'企业[：:]\s*([^\n\r，。；;\s]+)',
                r'任职于[：:]\s*([^\n\r，。；;\s]+)',
                r'工作单位[：:]\s*([^\n\r，。；;\s]+)',
                r'([\u4e00-\u9fa5]{2,}(公司|集团|企业|有限公司|科技|信息|软件|数据|互联网|网络|电子|通信))',
            ]
            
            for pattern in company_patterns:
                match = re.search(pattern, block)
                if match:
                    if match.lastindex:
                        exp.company = match.group(1).strip()
                    else:
                        exp.company = match.group().strip()
                    break
            
            position_patterns = [
                r'职位[：:]\s*([^\n\r，。；;\n]+)',
                r'岗位[：:]\s*([^\n\r，。；;\n]+)',
                r'担任[：:]\s*([^\n\r，。；;\n]+)',
                r'职务[：:]\s*([^\n\r，。；;\n]+)',
                r'([\u4e00-\u9fa5]{2,6}(工程师|设计师|经理|主管|总监|专员|顾问|分析师|开发|测试|产品))',
            ]
            
            for pattern in position_patterns:
                match = re.search(pattern, block)
                if match:
                    if match.lastindex:
                        exp.position = match.group(1).strip()
                    else:
                        exp.position = match.group().strip()
                    break
            
            desc_patterns = [
                r'工作内容[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
                r'工作职责[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
                r'工作描述[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
                r'职责描述[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
                r'主要职责[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
            ]
            
            for pattern in desc_patterns:
                match = re.search(pattern, block)
                if match:
                    exp.description = match.group(1).strip()
                    break
            
            if not exp.description:
                lines = [l.strip() for l in block.split('\n') if l.strip()]
                if lines:
                    start_idx = 0
                    for i, line in enumerate(lines):
                        if re.search(r'(工作内容|工作职责|工作描述|主要职责|项目)', line):
                            start_idx = i
                            break
                    if start_idx < len(lines) - 1:
                        exp.description = '\n'.join(lines[start_idx + 1:])
            
            if exp.company or exp.position or exp.start_date:
                experiences.append(exp)
        
        return experiences
    
    def _split_work_fuzzy(self, text: str) -> List[str]:
        blocks = []
        
        company_keywords = [
            '公司', '集团', '企业', '有限公司', '科技', '信息', '软件',
            '互联网', '网络', '电子', '通信', '数据'
        ]
        
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        
        current_block = ''
        in_work = False
        
        for line in lines:
            has_company = any(kw in line for kw in company_keywords)
            has_position = re.search(r'(工程师|设计师|经理|主管|总监|专员|顾问|分析师|开发|测试|产品)', line)
            has_date = re.search(r'\d{4}[\s年./-]*\d{0,2}', line)
            
            if has_company or has_position or has_date:
                if current_block:
                    blocks.append(current_block.strip())
                current_block = line + '\n'
                in_work = True
            elif in_work:
                current_block += line + '\n'
        
        if current_block:
            blocks.append(current_block.strip())
        
        if not blocks:
            blocks = [text.strip()]
        
        return blocks
    
    def _parse_project_section(self, section_text: str) -> List[ProjectExperience]:
        experiences = []
        
        project_blocks = self._split_experience_blocks(section_text)
        
        for block in project_blocks:
            exp = ProjectExperience()
            
            date_range = self._extract_date_range(block)
            if date_range:
                exp.start_date = date_range[0]
                exp.end_date = date_range[1]
            
            name_patterns = [
                r'项目名称[：:]\s*([^\n\r，。；;]+)',
                r'项目名[：:]\s*([^\n\r，。；;]+)',
            ]
            
            for pattern in name_patterns:
                match = re.search(pattern, block)
                if match:
                    exp.project_name = match.group(1).strip()
                    break
            
            role_patterns = [
                r'项目角色[：:]\s*([^\n\r，。；;]+)',
                r'担任角色[：:]\s*([^\n\r，。；;]+)',
                r'职责[：:]\s*([^\n\r，。；;]+)',
            ]
            
            for pattern in role_patterns:
                match = re.search(pattern, block)
                if match:
                    exp.role = match.group(1).strip()
                    break
            
            desc_patterns = [
                r'项目描述[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
                r'项目介绍[：:]\s*([\s\S]+?)(?=\n\s*\n|$)',
            ]
            
            for pattern in desc_patterns:
                match = re.search(pattern, block)
                if match:
                    exp.description = match.group(1).strip()
                    break
            
            tech_pattern = r'技术栈[：:]\s*([^\n\r]+)'
            match = re.search(tech_pattern, block)
            if match:
                techs = re.split(r'[,，、\s]+', match.group(1).strip())
                exp.technologies = [t for t in techs if t]
            
            if exp.project_name or exp.role:
                experiences.append(exp)
        
        return experiences
    
    def _split_experience_blocks(self, text: str) -> List[str]:
        blocks = []
        
        date_pattern = r'\d{4}[\s年./-]*\d{0,2}[\s月./-]*'
        patterns = [
            rf'({date_pattern}[\s至~-]+{date_pattern})',
            rf'(^|\n)({date_pattern})',
        ]
        
        for pattern in patterns:
            matches = list(re.finditer(pattern, text))
            if len(matches) >= 1:
                positions = [m.start() for m in matches]
                positions.append(len(text))
                
                for i in range(len(positions) - 1):
                    block = text[positions[i]:positions[i+1]].strip()
                    if block:
                        blocks.append(block)
                break
        
        if not blocks and text.strip():
            blocks = [text.strip()]
        
        return blocks
    
    def _extract_date_range(self, text: str) -> Optional[Tuple[Optional[datetime], Optional[datetime]]]:
        patterns = [
            r'(\d{4})[年./-]?(\d{0,2})[月]?[\s至~-]+(\d{4})[年./-]?(\d{0,2})[月]?(至今)?',
            r'(\d{4})[\s年./-]*(\d{0,2})[\s月./-]*[\s至~-]+(\d{4})[\s年./-]*(\d{0,2})[\s月]?(至今)?',
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                start_year = int(match.group(1))
                start_month = int(match.group(2)) if match.group(2) else 1
                end_year = int(match.group(3))
                end_month = int(match.group(4)) if match.group(4) else 1
                is_present = match.group(5) is not None
                
                try:
                    start_date = datetime(start_year, start_month, 1)
                    end_date = None if is_present else datetime(end_year, end_month, 1)
                    return (start_date, end_date)
                except ValueError:
                    pass
        
        return None
    
    def _extract_skills(self, text: str) -> List[str]:
        skills = []
        
        skill_patterns = [
            r'技能[：:]\s*([^\n\r]+)',
            r'专业技能[：:]\s*([^\n\r]+)',
            r'掌握技能[：:]\s*([^\n\r]+)',
            r'技术栈[：:]\s*([^\n\r]+)',
            r'技术能力[：:]\s*([^\n\r]+)',
            r'擅长领域[：:]\s*([^\n\r]+)',
            r'核心技能[：:]\s*([^\n\r]+)',
            r'熟练技能[：:]\s*([^\n\r]+)',
            r'技能特长[：:]\s*([^\n\r]+)',
            r'Skills[：:\s]+([^\n\r]+)',
            r'Technical[：:\s]+([^\n\r]+)',
        ]
        
        for pattern in skill_patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                skill_text = match.group(1)
                skill_list = re.split(r'[,，、\s;；·]+', skill_text)
                skills.extend([s.strip() for s in skill_list if s.strip() and len(s.strip()) > 1])
        
        common_skills = [
            'Python', 'Java', 'C++', 'C#', 'JavaScript', 'TypeScript', 'Go', 'Rust', 'PHP',
            'React', 'Vue', 'Angular', 'Node.js', 'Django', 'Flask', 'Spring', 'SpringBoot',
            'MySQL', 'PostgreSQL', 'Oracle', 'MongoDB', 'Redis', 'Elasticsearch',
            'Docker', 'Kubernetes', 'AWS', 'Linux', 'Git', 'SVN',
            '机器学习', '深度学习', '人工智能', '数据挖掘', '数据分析',
            'UI设计', 'UX设计', '产品经理', '项目管理',
        ]
        
        for skill in common_skills:
            if skill in text and skill not in skills:
                skills.append(skill)
        
        return list(set(skills))


class PDFParser(BaseParser):
    def __init__(self):
        try:
            import pdfplumber
            self.pdfplumber = pdfplumber
            self.available = True
        except ImportError:
            self.available = False
    
    def parse(self, file_path: str) -> ResumeData:
        if not self.available:
            raise ImportError("pdfplumber is not installed. Install it with: pip install pdfplumber")
        
        text = self._extract_text(file_path)
        return self._parse_text(text)
    
    def _extract_text(self, file_path: str) -> str:
        text_parts = []
        
        with self.pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        
        return '\n'.join(text_parts)
    
    def _parse_text(self, text: str) -> ResumeData:
        resume_data = ResumeData()
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        
        resume_data.raw_data['full_text'] = text
        
        resume_data.personal_info.name = self._extract_name(text, lines)
        resume_data.personal_info.phone = self._extract_phone(text)
        resume_data.personal_info.email = self._extract_email(text)
        resume_data.personal_info.gender = self._extract_gender(text)
        
        intro_patterns = [
            r'自我介绍[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
            r'个人简介[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
            r'自我评价[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
        ]
        
        for pattern in intro_patterns:
            match = re.search(pattern, text)
            if match:
                resume_data.personal_info.self_introduction = match.group(1).strip()
                break
        
        sections = self._split_by_sections(text)
        
        section_mappings = {
            '教育背景': 'education',
            '教育经历': 'education',
            '学习经历': 'education',
            '工作经历': 'work',
            '工作经验': 'work',
            '职业经历': 'work',
            '项目经历': 'project',
            '项目经验': 'project',
        }
        
        for section_name, section_content in sections.items():
            mapped_type = section_mappings.get(section_name)
            
            if mapped_type == 'education':
                resume_data.education_experiences.extend(
                    self._parse_education_section(section_content)
                )
            elif mapped_type == 'work':
                resume_data.work_experiences.extend(
                    self._parse_work_section(section_content)
                )
            elif mapped_type == 'project':
                resume_data.project_experiences.extend(
                    self._parse_project_section(section_content)
                )
        
        resume_data.skills = self._extract_skills(text)
        
        return resume_data


class WordParser(BaseParser):
    def __init__(self):
        try:
            from docx import Document
            self.Document = Document
            self.available = True
        except ImportError:
            self.available = False
    
    def parse(self, file_path: str) -> ResumeData:
        if not self.available:
            raise ImportError("python-docx is not installed. Install it with: pip install python-docx")
        
        text = self._extract_text(file_path)
        return self._parse_text(text)
    
    def _extract_text(self, file_path: str) -> str:
        doc = self.Document(file_path)
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(' '.join(row_text))
        
        return '\n'.join(text_parts)
    
    def _parse_text(self, text: str) -> ResumeData:
        resume_data = ResumeData()
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        
        resume_data.raw_data['full_text'] = text
        
        resume_data.personal_info.name = self._extract_name(text, lines)
        resume_data.personal_info.phone = self._extract_phone(text)
        resume_data.personal_info.email = self._extract_email(text)
        resume_data.personal_info.gender = self._extract_gender(text)
        
        intro_patterns = [
            r'自我介绍[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
            r'个人简介[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
            r'自我评价[：:]\s*([\s\S]+?)(?=\n\s*[【■◆●]|$)',
        ]
        
        for pattern in intro_patterns:
            match = re.search(pattern, text)
            if match:
                resume_data.personal_info.self_introduction = match.group(1).strip()
                break
        
        sections = self._split_by_sections(text)
        
        section_mappings = {
            '教育背景': 'education',
            '教育经历': 'education',
            '学习经历': 'education',
            '工作经历': 'work',
            '工作经验': 'work',
            '职业经历': 'work',
            '项目经历': 'project',
            '项目经验': 'project',
        }
        
        for section_name, section_content in sections.items():
            mapped_type = section_mappings.get(section_name)
            
            if mapped_type == 'education':
                resume_data.education_experiences.extend(
                    self._parse_education_section(section_content)
                )
            elif mapped_type == 'work':
                resume_data.work_experiences.extend(
                    self._parse_work_section(section_content)
                )
            elif mapped_type == 'project':
                resume_data.project_experiences.extend(
                    self._parse_project_section(section_content)
                )
        
        resume_data.skills = self._extract_skills(text)
        
        return resume_data


class ResumeParser:
    def __init__(self):
        self._pdf_parser = PDFParser()
        self._word_parser = WordParser()
    
    def parse(self, file_path: str) -> ResumeData:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume file not found: {file_path}")
        
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            return self._pdf_parser.parse(file_path)
        elif ext in ['.docx', '.doc']:
            if ext == '.doc':
                raise ValueError(".doc format is not supported. Please convert to .docx first.")
            return self._word_parser.parse(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: .pdf, .docx")
    
    def is_available(self, file_path: str) -> bool:
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            return self._pdf_parser.available
        elif ext in ['.docx', '.doc']:
            return self._word_parser.available
        
        return False
